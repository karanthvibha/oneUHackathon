import io
import json
import os
import re

from flask import Flask, jsonify, request
from flask_cors import CORS

try:  # pragma: no cover - dependency optional at import time
    import boto3
    from botocore.exceptions import BotoCoreError, ClientError
except ImportError:  # pragma: no cover
    boto3 = None
    BotoCoreError = ClientError = None


def _load_env():
    """Load environment variables from backend/.env if present (dev convenience)."""
    try:
        from dotenv import load_dotenv

        load_dotenv()
    except ImportError:
        pass


def _read_bool(value):
    return str(value).strip().lower() in {"1", "true", "yes", "on"}


def _decode_text(raw: bytes) -> str:
    for encoding in ("utf-8", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _extract_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        return ""
    try:
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages).strip()
    except Exception:
        return ""


def extract_text_from_upload(file_storage) -> str:
    filename = (file_storage.filename or "").lower()
    data = file_storage.read()
    if not data:
        return ""

    ext = filename.rsplit(".", 1)[-1] if "." in filename else ""

    if ext in {"txt", "md", "csv", "rtf", "html", "htm"}:
        return _decode_text(data).strip()

    if ext == "pdf":
        return _extract_pdf(data)

    if ext in {"doc", "docx"}:
        try:
            from docx import Document

            doc = Document(io.BytesIO(data))
            return "\n".join(p.text for p in doc.paragraphs).strip()
        except ImportError:
            return ""

    if ext in {"ppt", "pptx", "key"}:
        try:
            from pptx import Presentation

            prs = Presentation(io.BytesIO(data))
            chunks = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if getattr(shape, "text", None):
                        chunks.append(shape.text)
            return "\n".join(chunks).strip()
        except ImportError:
            return ""

    # Fallback: try to read the file as UTF-8 text.
    return _decode_text(data).strip()


def _build_notes_prompt(file_name: str, file_text: str, existing_notes, folders) -> str:
    notes_json = json.dumps(existing_notes, ensure_ascii=False)
    folders_json = json.dumps(folders, ensure_ascii=False)
    return (
        "You are a study assistant that converts source material into concise, well-structured Smart Notes.\n\n"
        f"Source file name: {file_name}\n\n"
        "Source file content:\n"
        f"{file_text}\n\n"
        f"Existing note folders: {folders_json}\n\n"
        f"Existing notes (title and body): {notes_json}\n\n"
        "Instructions:\n"
        "1. Decide whether any existing note is a good thematic fit for this material.\n"
        "2. If a suitable existing note is found, target that note and write the FULL updated note body "
        "that merges the new material into the existing note, keeping it clean and useful.\n"
        "3. If no suitable existing note is found, invent a concise but descriptive note title and write "
        "a fresh formatted notes body for the material.\n"
        "4. Format the body as study notes using short paragraphs and bullet points where helpful.\n\n"
        "Return ONLY valid JSON with exactly this shape (and no other text):\n"
        '{"targetTitle": "<the note title>", "body": "<the full formatted note body>"}\n'
    )


def _parse_notes_response(text: str):
    if not text:
        return None
    match = re.search(r"\{.*\}", text, re.S)
    candidate = match.group(0) if match else text
    try:
        data = json.loads(candidate)
    except json.JSONDecodeError:
        return None

    target_title = str(data.get("targetTitle", "")).strip()
    body = str(data.get("body", "")).strip()
    if not body:
        return None
    return {"targetTitle": target_title or "Imported notes", "body": body}


def create_app():
    _load_env()
    app = Flask(__name__)

    allowed_origins = os.getenv("CORS_ORIGINS", "http://localhost:5173")
    CORS(
        app,
        resources={r"/api/*": {"origins": [o.strip() for o in allowed_origins.split(",") if o.strip()]}},
    )

    def credentials_from_env():
        access_key = os.getenv("AWS_ACCESS_KEY_ID") or os.getenv("BEDROCK_ACCESS_KEY_ID")
        secret_key = os.getenv("AWS_SECRET_ACCESS_KEY") or os.getenv("BEDROCK_SECRET_ACCESS_KEY")
        session_token = os.getenv("AWS_SESSION_TOKEN") or os.getenv("BEDROCK_SESSION_TOKEN")
        return access_key, secret_key, session_token

    def region():
        return os.getenv("AWS_REGION") or os.getenv("BEDROCK_REGION") or "us-west-2"

    def default_model_id():
        return os.getenv("BEDROCK_MODEL_ID", "anthropic.claude-3-5-sonnet-20241022-v2:0")

    def get_bedrock_client():
        if boto3 is None:
            raise RuntimeError("boto3 is not installed. Run: pip install -r requirements.txt")

        access_key, secret_key, session_token = credentials_from_env()
        kwargs = {"region_name": region()}
        if access_key and secret_key:
            kwargs["aws_access_key_id"] = access_key
            kwargs["aws_secret_access_key"] = secret_key
        if session_token:
            kwargs["aws_session_token"] = session_token

        return boto3.client("bedrock-runtime", **kwargs)

    def bedrock_error_response(exc):
        if ClientError and isinstance(exc, ClientError):
            code = exc.response.get("Error", {}).get("Code", "ClientError")
            message = exc.response.get("Error", {}).get("Message", str(exc))
        else:
            code = "BedrockError"
            message = str(exc)
        return jsonify({"error": message, "code": code}), 502

    @app.get("/api/health")
    def health():
        return jsonify({"status": "ok"})

    @app.get("/api/bedrock/status")
    def bedrock_status():
        access_key, secret_key, _ = credentials_from_env()
        env_secrets_set = bool(access_key and secret_key)
        return jsonify(
            {
                "configured": env_secrets_set,
                "region": region(),
                "modelId": default_model_id(),
                "usingIamRole": not env_secrets_set,
            }
        )

    @app.post("/api/bedrock/invoke")
    def bedrock_invoke():
        if boto3 is None:
            return jsonify({"error": "Server missing the boto3 dependency."}), 500

        data = request.get_json(silent=True) or {}
        model_id = data.get("modelId") or default_model_id()
        messages = data.get("messages")
        if not isinstance(messages, list) or not messages:
            return jsonify({"error": "A non-empty 'messages' list is required."}), 400

        inference_config = data.get("inferenceConfig") or {}
        try:
            temperature = float(inference_config.get("temperature", 0.7))
            max_tokens = int(inference_config.get("maxTokens", 1024))
        except (TypeError, ValueError):
            temperature, max_tokens = 0.7, 1024

        try:
            client = get_bedrock_client()
            response = client.converse(
                modelId=model_id,
                messages=messages,
                inferenceConfig={"temperature": temperature, "maxTokens": max_tokens},
            )
        except (ClientError, BotoCoreError) as exc:
            return bedrock_error_response(exc)
        except Exception as exc:  # noqa: BLE001
            return jsonify({"error": str(exc)}), 500

        output = response.get("output", {})
        message = output.get("message", {})
        content = message.get("content", [])
        text = "".join(block.get("text", "") for block in content if block.get("text"))
        return jsonify({"text": text, "modelId": model_id})

    @app.post("/api/bedrock/integrate-notes")
    def integrate_notes():
        if boto3 is None:
            return jsonify({"error": "Server missing the boto3 dependency."}), 500

        if "file" not in request.files:
            return jsonify({"error": "A multipart 'file' field is required."}), 400

        file_storage = request.files["file"]
        file_name = file_storage.filename or "material"
        file_text = extract_text_from_upload(file_storage)
        if not file_text:
            return jsonify(
                {"error": "Could not extract readable text from this file. Try a text-based or PDF file."}
            ), 400

        try:
            existing_notes = json.loads(request.form.get("existingNotesJson", "[]"))
            folders = json.loads(request.form.get("foldersJson", "[]"))
        except json.JSONDecodeError:
            return jsonify({"error": "existingNotesJson and foldersJson must be valid JSON."}), 400

        model_id = request.form.get("modelId") or default_model_id()
        prompt = _build_notes_prompt(file_name, file_text[:12000], existing_notes, folders)

        try:
            client = get_bedrock_client()
            response = client.converse(
                modelId=model_id,
                messages=[{"role": "user", "content": [{"text": prompt}]}],
                inferenceConfig={"temperature": 0.2, "maxTokens": 2048},
            )
        except (ClientError, BotoCoreError) as exc:
            return bedrock_error_response(exc)
        except Exception as exc:  # noqa: BLE001
            return jsonify({"error": str(exc)}), 500

        output = response.get("output", {})
        message = output.get("message", {})
        content = message.get("content", [])
        llm_text = "".join(block.get("text", "") for block in content if block.get("text"))

        result = _parse_notes_response(llm_text)
        if not result:
            return jsonify({"error": "Model returned an unparseable response.", "raw": llm_text}), 502

        return jsonify(result)

    return app


app = create_app()


if __name__ == "__main__":
    port = int(os.getenv("PORT", "5001"))
    debug = _read_bool(os.getenv("FLASK_DEBUG", "0"))
    app.run(host="0.0.0.0", port=port, debug=debug)